"""Daily report: markdown summary + a spreadsheet you can open on your phone (reports/Paper Bot Live.xlsx)."""
import os, datetime as dt
import pandas as pd

ACCTS = [200, 250, 350, 400, 550, 1000]


def _read(path):
    return pd.read_csv(path) if os.path.exists(path) else pd.DataFrame()


def daily(base, day):
    os.makedirs(f"{base}/reports", exist_ok=True)
    j, a, s = _read(f"{base}/journal.csv"), _read(f"{base}/alpaca_journal.csv"), _read(f"{base}/signals.csv")
    for df in (j, a, s):
        if len(df): df["date"] = df["date"].astype(str)
    d = str(day); jt = j[j.date == d] if len(j) else j; st = s[s.date == d] if len(s) else s
    lines = [f"# Paper trading report {d}", ""]
    lines += ["## Today's trades", ""]
    if len(jt):
        lines += ["| Account | Strategy | Ticker | Entry | Exit | Why | P&L | R |", "|---|---|---|---|---|---|---|---|"]
        lines += [f"| ${r.account} | {r.strategy}{' (test)' if getattr(r, 'test', 0) == 1 else ''} | {r.ticker} | {r.entry} | {r.exit} | {r.exit_reason} | ${r.pnl:.2f} | {r.R:+.2f} |" for r in jt.itertuples()]
        if "test" in jt and (jt.test == 1).any():
            lines += ["", "(test) = strategy the backtest switched off; traded on paper only to measure live fills and timing."]
        lines += ["", f"**Day total: ${jt.pnl.sum():.2f}** across {len(jt)} trades (goal: $20/day)", ""]
    else:
        lines += ["No trades closed today.", ""]
    if len(a) and len(a[a.date == d]):
        at = a[a.date == d]
        lines += ["## Simulator vs Alpaca fills", "", "| Account | Ticker | Sim P&L | Alpaca P&L | Status |", "|---|---|---|---|---|"]
        lines += [f"| ${r.account} | {r.ticker} | ${r.sim_pnl} | {'' if pd.isna(r.alp_pnl) else '$' + str(r.alp_pnl)} | {r.status} |" for r in at.itertuples()]
        lines += [""]
    if len(st):
        lab = st.dropna(subset=["R"])
        lines += ["## Every signal today (taken or not)", "", f"{len(st)} signals, {len(lab)} labeled, average {lab.R.mean():+.2f}R" if len(lab) else f"{len(st)} signals", ""]
        if len(lab):
            g = lab.groupby("strategy").R.agg(["count", "mean"]).round(2).sort_values("mean", ascending=False)
            lines += ["| Strategy | Signals | Avg R |", "|---|---|---|"] + [f"| {k} | {int(r['count'])} | {r['mean']:+.2f} |" for k, r in g.iterrows()] + [""]
    if len(j):
        tot = j.groupby("account").pnl.sum()
        lines += ["## Running totals", "", "| Account | Trades | P&L | Value |", "|---|---|---|---|"]
        for acct in ACCTS:
            n = int((j.account == acct).sum()); p = float(tot.get(acct, 0.0))
            lines += [f"| ${acct} | {n} | ${p:.2f} | ${acct + p:.2f} |"]
        days = j.date.nunique()
        lines += ["", f"All accounts: ${j.pnl.sum():.2f} over {days} trading days = **${j.pnl.sum() / max(days, 1):.2f}/day average**"]
    open(f"{base}/reports/{d}.md", "w").write("\n".join(lines) + "\n")
    workbook(base, j, a, s)


def workbook(base, j, a, s):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    F = lambda **k: Font(name="Arial", **k); H = PatternFill("solid", fgColor="1F2937")
    wb = Workbook(); ws = wb.active; ws.title = "Accounts"
    ws["A1"] = "Paper Bot Live: accounts"; ws["A1"].font = F(bold=True, size=14)
    ws["A2"] = "Rebuilt by the bot after every session from journal.csv. Totals are formulas over the Journal tab."; ws["A2"].font = F(italic=True, color="6B7280")
    hdr = ["Account ($)", "Trades", "Wins", "Win rate", "P&L ($)", "R total", "Avg R", "Value ($)"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(4, i, h); c.font = F(bold=True, color="FFFFFF"); c.fill = H
    for r, acct in enumerate(ACCTS, 5):
        ws.cell(r, 1, acct)
        ws.cell(r, 2, f"=COUNTIF(Journal!$B:$B,A{r})"); ws.cell(r, 3, f'=COUNTIFS(Journal!$B:$B,A{r},Journal!$L:$L,">0")')
        ws.cell(r, 4, f"=IFERROR(C{r}/B{r},0)"); ws.cell(r, 5, f"=SUMIF(Journal!$B:$B,A{r},Journal!$L:$L)")
        ws.cell(r, 6, f"=SUMIF(Journal!$B:$B,A{r},Journal!$M:$M)"); ws.cell(r, 7, f"=IFERROR(F{r}/B{r},0)"); ws.cell(r, 8, f"=A{r}+E{r}")
        for c, fm in [(1, "$#,##0"), (4, "0%"), (5, "$#,##0.00;($#,##0.00);-"), (6, "0.00"), (7, "0.00"), (8, "$#,##0.00")]: ws.cell(r, c).number_format = fm
    ws.cell(11, 1, "Total").font = F(bold=True)
    for c, f, fm in [(2, "=SUM(B5:B10)", None), (3, "=SUM(C5:C10)", None), (4, "=IFERROR(C11/B11,0)", "0%"), (5, "=SUM(E5:E10)", "$#,##0.00;($#,##0.00);-"),
                     (6, "=SUM(F5:F10)", "0.00"), (7, "=IFERROR(F11/B11,0)", "0.00"), (8, "=SUM(H5:H10)", "$#,##0.00")]:
        ws.cell(11, c, f).font = F(bold=True)
        if fm: ws.cell(11, c).number_format = fm
    ws["A13"] = "Trading days"; ws["B13"] = "=SUMPRODUCT(1/COUNTIF(Journal!A2:A5000,Journal!A2:A5000&\"\"))-1"
    ws["A14"] = "Average $/day"; ws["B14"] = "=IFERROR(E11/B13,0)"; ws["B14"].number_format = "$#,##0.00"
    ws["A15"] = "Goal $/day"; ws["B15"] = 20; ws["B15"].font = F(color="0000FF"); ws["B15"].number_format = "$#,##0.00"
    for col, w in zip("ABCDEFGH", [14, 9, 8, 10, 12, 9, 9, 12]): ws.column_dimensions[col].width = w
    for name, df in [("Journal", j), ("Signals", s), ("Alpaca fills", a)]:
        sh = wb.create_sheet(name)
        if not len(df): sh["A1"] = "No data yet"; continue
        cols = list(df.columns)
        for i, h in enumerate(cols, 1):
            c = sh.cell(1, i, h); c.font = F(bold=True, color="FFFFFF"); c.fill = H
        for r, row in enumerate(df.itertuples(index=False), 2):
            for i, v in enumerate(row, 1):
                sh.cell(r, i, None if (isinstance(v, float) and pd.isna(v)) else v)
        sh.freeze_panes = "A2"
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if c.font is None or c.font.name != "Arial": c.font = F(bold=c.font.b if c.font else False, color=c.font.color if c.font else None)
    wb.save(f"{base}/reports/Paper Bot Live.xlsx")
